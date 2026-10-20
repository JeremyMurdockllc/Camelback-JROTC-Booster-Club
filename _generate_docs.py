# -*- coding: utf-8 -*-
"""Generate printable PDF playbook and Word bylaws/roster for the Camelback booster kit."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

KIT = Path(r"C:\Users\jerem\Desktop\Camelback Booster Club Startup Kit")
PDF_CREATE = Path(r"C:\Users\jerem\AppData\Local\hermes\skills\productivity\pdf\scripts\pdf_create.py")
SPEC = KIT / "_playbook_spec.json"
PDF = KIT / "Camelback_Booster_Club_Startup_Playbook.pdf"
BYLAWS = KIT / "Camelback_Booster_Club_Bylaws_DRAFT.docx"
ROSTER = KIT / "Camelback_Booster_Officer_Roster_Worksheet.docx"


def p(text: str) -> dict:
    return {"type": "paragraph", "text": text}


def h(text: str, level: int = 1) -> dict:
    return {"type": "heading", "text": text, "level": level}


def tbl(rows: list[list[str]], header: bool = True) -> dict:
    return {"type": "table", "rows": rows, "header": header}


def br() -> dict:
    return {"type": "pagebreak"}


def build_spec() -> dict:
    elements = [
        h("Camelback High School Booster Club Startup Playbook"),
        p("Phoenix Union High School District No. 210. Working packet for parent organizers. Not legal, tax, or district advice. Confirm every filing with PXU, the Arizona Corporation Commission, the IRS, and (when needed) an Arizona nonprofit attorney or CPA. Fees and forms change."),
        p("Companion files: open index.html in a browser to fill officer names (saved on that computer), export a CSV/JSON roster, and refresh the bylaws draft. This PDF is the printable checklist."),
        h("Do this in this order", 2),
        p("1. Meet the Camelback principal (and Athletic Director or program sponsor). Get written permission to organize. Do not fundraise yet. PXU KJA-OP-1: unapproved boosters may not operate or use facilities."),
        p("2. Recruit President, Vice President, Secretary, Treasurer. District employees may advise only — they cannot be officers or check signers."),
        p("3. Choose a legal name the principal approves. Do not put Phoenix Union, PUHSD, or the school name in the title or on the bank account. Spartan mascot is allowed with permission."),
        p("4. Adopt bylaws (Word draft in this folder) and a one-paragraph purpose. PXU requires a constitution/bylaws. IRS expects 501(c)(3) clauses."),
        p("5. File Arizona nonprofit Articles of Incorporation on stock Form C011 (ACC default language). Not an LLC. Not Articles of Organization. Cover sheet + C003 + statutory-agent acceptance. $40."),
        p("6. Wait for ACC approval. For Phoenix / Maricopa, no newspaper ad. ACC posts the approval in its public database (A.R.S. 10-3203(D)(2) and 10-130). Old C011i PDFs that still say “publish in a newspaper” are outdated."),
        p("7. Get a free EIN from IRS.gov (Form SS-4). Print the confirmation. Then open a two-signature bank account in the corporate name. Never use the school EIN or school letterhead."),
        p("8. File ACC Form C015 (Articles of Amendment – Nonprofit) adding IRS purpose, inurement, politics, and dissolution clauses. $25. Do this before Pay.gov."),
        p("9. Apply for 501(c)(3) on Pay.gov (usually Form 1023-EZ, $275 user fee) only after the C015 amendment is approved. Donations are not deductible until IRS recognition."),
        p("10. Submit two signed PXU Booster Club Applications plus attachments to the principal (then District finance). Re-apply every year. Year-end packet due to the principal by August 31."),
        h("1. Phoenix Union approval", 2),
        p("Policy KJA / procedure KJA-OP-1. Only an approved booster may use PXU facilities. Approval is annual. The school keeps the original application; copies go to the District Comptroller / Executive Director of Finance."),
        h("Camelback contacts (confirm before you email)", 3),
        tbl(
            [
                ["Role", "Name", "Contact"],
                ["Athletic Director", "Daniel B. Eaton", "deaton@phoenixunion.org / (602) 764-7007"],
                ["Athletics Administrative Assistant", "Brieanna Garcia", "bgarcia1@phoenixunion.org"],
                ["Campus", "Camelback High School", "4612 N. 28th Street, Phoenix, AZ 85016"],
                ["PXU Student Activities", "District office", "studentactivities@phoenixunion.org"],
                ["Principal", "Confirm at the first meeting", "Signature required on the application"],
                ["Program sponsor", "Coach / JROTC instructor / certified staff", "Advisor only. Not an officer. Not a check signer."],
            ]
        ),
        p("Official packet: https://www.pxu.org/documents/departments/bookstore/booster-club/639148  and  https://www.pxu.org/page/bookstore/"),
        p("Download from that page: Booster Club Application; Formation and Operation of Booster Club (KJA-OP-1); Establishing a Non-Profit Organization; sample financial statement and bank reconciliation. The older Establishing a Non-Profit handout still cites outdated IRS dollar amounts (it says Form 990-N at $25,000; the current IRS threshold is normally $50,000) and old ACC web addresses. Use current IRS/ACC links in section 4–7."),
        h("PXU application attachments (KJA-OP-1 EB)", 3),
        tbl(
            [
                ["Item", "New formal corporation", "Notes"],
                ["Application (two signed copies)", "Yes", "Officers: name, address, phone, email for President, VP, Secretary, Treasurer"],
                ["Articles of Incorporation", "Yes (new formal corp)", "Stamped ACC copy"],
                ["Constitution and bylaws", "Yes", "See Word draft in this folder"],
                ["Proposed activities, events, budget", "Yes", ""],
                ["Anticipated maximum fiscal-year donations", "Yes", "Governing Board pre-approves cash donations up to this figure"],
                ["Bank statement + authorized withdrawers", "When account exists", "Two signers"],
                ["IRS 501(c)(3) determination letter", "When received", "You may incorporate first and add the letter later"],
                ["Prior-year treasurer / income statement", "Renewals", ""],
                ["ACC annual report (prior year)", "Renewals", ""],
            ]
        ),
        p("Hard PXU rules: no individual student fundraising accounts; students pay required trip/participation fees to the campus bookstore; two signatures on every check; do not use the district EIN; do not use school letterhead; fundraising needs prior written principal approval; boosters should not sign student-travel contracts; employees of PXU should not be officers or account signers."),
        br(),
        h("2. Names and positions you must collect", 2),
        p("Fill the browser kit (index.html) or the Word roster worksheet. Do not put Social Security numbers in either file. The EIN responsible party enters an SSN only on IRS.gov."),
        p("President, VP, Secretary, Treasurer — PXU application, bank, 1023-EZ, bylaws signatures. Collect legal names, mailing addresses, phones, emails."),
        p("Directors — aim for three people who are not related. Officers may also be directors. Needed for Arizona articles and IRS."),
        p("Incorporator(s) — sign Arizona articles under penalty of law. Usually a founding director."),
        p("Statutory agent — Arizona articles plus Form M002 acceptance. An Arizona individual or authorized entity with a physical AZ street address. The corporation cannot be its own agent."),
        p("EIN responsible party — a real person who controls the club. SSN/ITIN on the IRS site only."),
        p("Two check signers — usually President and Treasurer. Not a PXU employee."),
        p("Certified staff sponsor — PXU liaison, ex officio advisor only. Not an officer and not a signer."),
        p("Also lock: principal-approved legal name; one-paragraph mission; Arizona street address for known place of business; mailing address that is NOT the school mailroom (PXU will not receive, sort, or distribute booster mail); fiscal year end (this kit defaults to July 31 / school-year cycle)."),
        h("3. Officer roles (PXU rewrite of the Skyline 2/22/26 sheet)", 2),
        p("The Skyline High School JROTC Booster Club roles sheet from the February 22, 2026 meeting is a useful duty split. Mesa Public Schools forms, MPS Club Sponsor agreements, and “First Sergeant” titles do not apply at Camelback. Rewrite below is for Phoenix Union."),
        p("President — liaison to the principal. Presides; signs the annual PXU application; keeps the 3–4 person spend-approval channel; submits every fundraising request for prior written principal approval; signs contracts and checks with board approval."),
        p("Vice President — backup and second vote. Votes promptly on expenditures; acts as President when needed; recruits volunteers; helps on large events."),
        p("Treasurer — financial custodian. Deposits all money; two-signature checks only; bank reconciliation within 30 days with a second-officer review; archives approvals next to receipts; reports at meetings; year-end statement for the principal by August 31; IRS 990 series after exemption; coordinates receipt backup with the sponsor (sponsor is not a signer)."),
        p("Secretary — records officer. Minutes (including acceptance of treasurer reports); membership roll; custodian of bylaws, articles, EIN, IRS letter, and district approvals; matching approval logs to minutes; turn records over within 30 days of election; keep financial records five years."),
        p("Whole-board rule: spending waits for 3–4 yes votes (or a minuted board vote) before the Treasurer pays. No one — including the coach or JROTC instructor — spends club money alone."),
        br(),
        h("4. Arizona nonprofit corporation", 2),
        p("File Articles of Incorporation for a nonprofit corporation. Do not file Articles of Organization (those are for LLCs). Do not form an LLC. This kit’s path: stock Form C011 first, then Form C015 amendment with IRS language."),
        p("File online at the Arizona Business Center (replaced eCorp on January 12, 2026): https://arizonabusinesscenter.azcc.gov/  Forms: https://azcc.gov/corporations/forms/corporation-forms  Paper: ACC Corporations Division, Examination Section, 1300 W. Washington St., Phoenix, AZ 85007. (602) 542-3026. answers@azcc.gov"),
        tbl(
            [
                ["Document", "Form", "Regular fee"],
                ["Cover sheet", "ACC cover sheet", "—"],
                ["Articles of Incorporation – Nonprofit (ACC default)", "C011", "$40 (expedite +$35)"],
                ["Certificate of Disclosure", "C003 (required ACC form)", "—"],
                ["Statutory Agent Acceptance", "M002", "—"],
                ["Director / Incorporator extra pages", "C082 / C084 if needed", "—"],
            ]
        ),
        p("C011 meets Arizona law only and does not include IRS 501(c)(3) language. That is expected. Do not delay the first filing to draft custom articles. After ACC approval, file C015."),
        p("No newspaper publication for this Phoenix club. A.R.S. 10-3203(D) lets ACC post the approval in its public database when the known place of business is in a county over 800,000 people (Maricopa). Do not file bylaws with the ACC. File the Arizona corporation annual report every year (nonprofit annual report fee is typically $10 — confirm on the fee schedule) or the corporation can be administratively dissolved."),
        h("5. Federal EIN", 2),
        p("Free from the IRS. Anyone charging for an EIN is a scam. Apply only after the Arizona corporation exists. Applying too early also starts the three-year clock for automatic revocation if you later forget Form 990-N."),
        p("Online: https://www.irs.gov/businesses/small-businesses-self-employed/get-an-employer-identification-number  Form SS-4: https://www.irs.gov/pub/irs-pdf/fss4.pdf  Instructions: https://www.irs.gov/pub/irs-pdf/iss4.pdf"),
        p("Complete in one sitting (15-minute timeout). Print the confirmation immediately. Limit: one EIN per responsible party per day. Entity type: corporation / nonprofit. Legal name must match the articles. Responsible party is a real controlling individual with that person’s SSN or ITIN. Never use the Phoenix Union or Camelback High School EIN."),
        h("6. Amend the articles so they are 501(c)(3)-ready (C015)", 2),
        p("After ACC approves C011, file Articles of Amendment – Nonprofit (C015). That is the Arizona nonprofit-corporation amendment — not LLC Articles of Organization, and not LLC Form L015."),
        p("Why two steps: C011 is a fill-in ACC form examiners already know. Custom IRS paragraphs on the first filing can slow or bounce a volunteer packet. Extra cost is $25. Extra time is one more ACC cycle. Do not file Form 1023-EZ until this amendment is approved — 1023-EZ attests that the IRS clauses already exist."),
        p("Packet: cover sheet + C015 + the IRS-language attachment in this folder + board minutes or member vote. Regular fee $25 (expedite $60) — confirm on the ACC fee schedule."),
        p("Use the IRS suggested language from Publication 557, not a paraphrase: https://www.irs.gov/charities-non-profits/suggested-language-for-corporations-and-associations-per-publication-557"),
        p("Required clauses: (1) organized exclusively for charitable and educational 501(c)(3) purposes; (2) no inurement of net earnings to members or officers except reasonable compensation; (3) no substantial lobbying and no candidate campaigns; (4) on dissolution, leftover assets go only to 501(c)(3) purposes or to government for a public purpose — not back to parents."),
        p("Same notice rule as formation: for Phoenix / Maricopa, ACC database posting under A.R.S. 10-11006(B)(2) and 10-130 replaces a newspaper ad. No newspaper required."),
        br(),
        h("7. Federal 501(c)(3) application", 2),
        p("Arizona nonprofit status is not federal tax exemption. File within 27 months of formation if you want exemption retroactive to the incorporation date."),
        p("Form 1023-EZ (typical new booster): $275 user fee on Pay.gov. Complete the Eligibility Worksheet in the 1023-EZ instructions. Common ceilings: projected annual gross receipts $50,000 or less and assets $250,000 or less. Several automatic disqualifiers (school, church, hospital, foreign address, certain successors, LLCs, etc.)."),
        p("Full Form 1023: $600 user fee. Use this if the club is not eligible for 1023-EZ. File on Pay.gov (search 1023)."),
        p("1023-EZ about page: https://www.irs.gov/forms-pubs/about-form-1023-ez  Instructions + worksheet: https://www.irs.gov/pub/irs-pdf/i1023ez.pdf  Full 1023: https://www.irs.gov/forms-pubs/about-form-1023  Fees: https://www.irs.gov/charities-non-profits/form-1023-and-1023-ez-amount-of-user-fee  Pay.gov: https://www.pay.gov/public/home"),
        h("Have this folder before you file 1023-EZ (attestation — still keep the papers)", 3),
        tbl(
            [
                ["Document", "Why"],
                ["Stamped C011 plus stamped C015 with IRS clauses", "You attest the IRS clauses are already in them"],
                ["ACC approval in the public database", "No newspaper affidavit for Phoenix / Maricopa"],
                ["Signed bylaws / constitution", "PXU and IRS"],
                ["EIN confirmation letter", "Required on the form"],
                ["Officer/director list with mailing addresses", "1023-EZ and PXU"],
                ["Completed 1023-EZ Eligibility Worksheet", "Keep it; do not upload it"],
                ["Conflict-of-interest policy", "Strongly recommended"],
                ["Short activity narrative + first-year budget", "You will describe activities on the form"],
                ["Pay.gov login for an officer who may sign, plus the user fee", "Electronic filing only"],
            ]
        ),
        p("Full Form 1023 also uploads the organizing document, amendments, bylaws, and financials, and may need Form 2848 if a CPA or attorney will speak to the IRS."),
        p("After exemption, file a 990-series return every year. Form 990-N (e-Postcard) if gross receipts are normally $50,000 or less: https://www.irs.gov/charities-non-profits/annual-electronic-filing-requirement-for-small-exempt-organizations-form-990-n-e-postcard  Due the 15th day of the 5th month after fiscal year end (July 31 year-end → December 15). Miss three consecutive years and 501(c)(3) is automatically revoked."),
        p("Parent Booster USA and similar firms will file for a fee. Not required. They still need this roster and the PXU-approved name."),
        h("8. Bank, insurance, money", 2),
        p("Account in the corporate name with the club EIN. Two signatures on every check. No school or district words in the account title. No individual student accounts (that can destroy 501(c)(3) and create taxable income to a family). Two people count cash; deposit promptly; keep tallies and invoices. Student travel: club donates to the school; students pay required fees to the bookstore; boosters should not sign hotel/bus contracts. Cash donations above the amount on the PXU application need extra Governing Board approval; equipment donations always need Board approval. PXU Trust insurance may cover an approved booster when a school administrator co-sponsors the activity; many clubs still buy their own liability policy. Independent annual review of the books (need not be a CPA)."),
        br(),
        h("9. Bylaws notes", 2),
        p("A Word draft is in this folder: Camelback_Booster_Club_Bylaws_DRAFT.docx. It is adapted from the Skyline High School JROTC Booster Club bylaws (2025–2026) with PXU constitution items and IRS dissolution language added. Skyline officer names and the Mesa address are not used. Do not file bylaws with the ACC. Vote them in, sign, and keep with the PXU packet."),
        p("Material changes from the Skyline template: PXU employees cannot be officers or signers (Skyline’s bylaws allowed school employees as officers); two-signature checks and no individual student accounts; digital 3–4 member spend approval; year-end report to the principal by August 31; principal-approved name that does not use the school or PUHSD words."),
        h("10. Annual calendar", 2),
        tbl(
            [
                ["When", "What"],
                ["Before any event", "Principal’s written approval for fundraiser / activity / facility use"],
                ["Each month", "Bank reconciliation within 30 days, second-officer review"],
                ["Each general meeting", "Treasurer report; minutes; approval log vs. receipts"],
                ["By August 31", "PXU annual application + year-end financial report to the principal"],
                ["ACC anniversary", "Arizona corporation annual report"],
                ["15th day of 5th month after FYE", "IRS Form 990-N / 990-EZ / 990"],
                ["End of business year", "Elect officers; turn records over within 30 days"],
            ]
        ),
        h("Sources", 2),
        p("PXU Booster Club Resources: https://www.pxu.org/documents/departments/bookstore/booster-club/639148"),
        p("PXU Bookstore / Student Activities: https://www.pxu.org/page/bookstore/"),
        p("PXU KJA-OP-1 Formation and Operation of Booster Clubs; Booster Club Application (KJA-OP-1 EB); Establishing a Non-Profit Organization."),
        p("Camelback athletics contacts: https://www.pxu.org/o/chs/page/athletics/"),
        p("ACC corporation forms: https://azcc.gov/corporations/forms/corporation-forms"),
        p("ACC C011i instructions (nonprofit articles, $40 fee, publication, IRS-language warning). Arizona Business Center: https://arizonabusinesscenter.azcc.gov/"),
        p("IRS Get an EIN; Publication 557 suggested article language; Form 1023-EZ / 1023 about pages and user fees ($275 / $600); Form 990-N ($50,000 threshold)."),
        p("Skyline JROTC Booster Club bylaws (2025–2026) and 2/22/26 roles sheet — templates only; not Camelback governing documents."),
    ]
    return {
        "title": "Camelback High School Booster Club Startup Playbook",
        "author": "Camelback booster organizers",
        "page_size": "letter",
        "page_numbers": True,
        "elements": elements,
    }


def write_bylaws_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    def title(text: str) -> None:
        t = doc.add_paragraph(text)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.runs[0].bold = True
        t.runs[0].font.size = Pt(14)

    def hdg(text: str) -> None:
        p = doc.add_paragraph(text)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(12)

    def body(text: str) -> None:
        doc.add_paragraph(text)

    title("BYLAWS OF [LEGAL NAME]")
    body(
        "Draft template for board review. Do not file with the Arizona Corporation Commission. "
        "Adapted from the Skyline High School JROTC Booster Club bylaws (2025–2026) and rewritten "
        "for Phoenix Union High School District / Camelback High School. Replace every item in [brackets], "
        "vote the document in, and sign. This is not a finished legal instrument."
    )
    hdg("Article I — Name")
    body("1. The name of the organization is [LEGAL NAME], referred to as the Organization.")
    body("2. The mailing address is [PO BOX OR OFFICER ADDRESS — not the school mailroom].")
    body(
        "3. The Organization supports the [PROGRAM: JROTC / athletics / band / other] program at "
        "Camelback High School, Phoenix Union High School District No. 210. The Organization is independent "
        "of the District. It shall not use the District or school name in its title or bank accounts except as "
        "the principal has approved in writing. Use of the Spartan mascot in the title requires principal permission."
    )
    hdg("Article II — Purpose and Policies")
    body(
        "1. The Organization is organized exclusively for charitable and educational purposes within the meaning "
        "of section 501(c)(3) of the Internal Revenue Code, or the corresponding section of any future federal tax code, "
        "including to support students in the [PROGRAM] program at Camelback High School by raising funds and volunteer "
        "support, in cooperation with school administration, for activities, equipment, and educational experiences."
    )
    body("2. The Organization is noncommercial, nonsectarian, and nonpartisan.")
    body(
        "3. The Organization will work with school administration, faculty, and students to provide quality experiences "
        "for students, promote student welfare, keep members informed, bring home and school closer, and raise and "
        "distribute funds to support the fiscal needs of the [PROGRAM] program."
    )
    body(
        "4. The Organization will not, directly or indirectly, participate or intervene in any political campaign on "
        "behalf of or in opposition to any candidate for public office."
    )
    body(
        "5. No part of the net earnings of the Organization shall inure to the benefit of, or be distributable to, its "
        "members, officers, directors, or other private persons, except that the Organization may pay reasonable "
        "compensation for services rendered and make payments in furtherance of its purposes."
    )
    body("6. Records are open for member review and will be made available to the School or District upon request.")
    body(
        "7. The Organization will adhere to Phoenix Union High School District policies and operating procedure KJA-OP-1 "
        "(and any successor), including financial controls, conflict of interest, fundraising approval, and recordkeeping."
    )
    body(
        "8. Upon the dissolution of the Organization, assets shall be distributed for one or more exempt purposes within "
        "the meaning of section 501(c)(3) of the Internal Revenue Code, or the corresponding section of any future federal "
        "tax code, or shall be distributed to the federal government, or to a state or local government, for a public purpose, "
        "preferably for the benefit of the [PROGRAM] program at Camelback High School. Any such assets not so disposed of "
        "shall be disposed of by a Court of Competent Jurisdiction of the county in which the principal office of the "
        "Organization is then located, exclusively for such purposes. Assets shall not be distributed to members or private individuals."
    )
    hdg("Article III — Membership")
    body(
        "1. Parents, guardians, teachers, other employees, and community members willing to uphold these bylaws may enroll "
        "as members by registering name, address, telephone, and email."
    )
    body("2. Membership is without regard to race, ethnicity, religion, gender, or disability.")
    body("3. Annual enrollment is conducted, but persons may join at any time.")
    body(
        "4. Only enrolled members may vote or hold office. Members must be enrolled at least 30 days before an officer "
        "election to vote in that election."
    )
    body(
        "5. A parent’s membership or dues shall not be linked to expenditures on behalf of that parent’s student. The "
        "Organization shall not exist to benefit only the children of participating parents."
    )
    body("6. The Executive Board may terminate membership for cause, including misconduct at Organization events.")
    hdg("Article IV — Business Year")
    body(
        "The business year begins August 1 and ends July 31 of the following calendar year, unless the board adopts a "
        "different fiscal year consistent with IRS filings. A year-end financial report is due to the principal by August 31."
    )
    hdg("Article V — Officers and Executive Board")
    body(
        "1. Officers are President, Vice President, Secretary, and Treasurer, provided members are willing to serve. "
        "One person may serve as Vice President and Secretary. A committee chair may also serve as an officer."
    )
    body(
        "2. Employees of Phoenix Union High School District shall not serve as officers or as authorized signers on the "
        "Organization’s bank accounts. They may serve as ex officio, non-voting advisors."
    )
    body(
        "3. The Executive Board is the elected officers, committee chairs, and the school program sponsor(s) in an "
        "ex officio advisory capacity."
    )
    hdg("Article VI — Election and Removal")
    body("1. Officers serve a one-year term beginning at the close of the meeting at which they are elected.")
    body(
        "2. Officers are elected by secret ballot at the final general membership meeting of the business year, unless "
        "the Executive Board schedules another meeting. Uncontested offices may be filled by voice vote."
    )
    body("3. A member must appear in person to vote.")
    body(
        "4. A candidate may hold the same office no more than two consecutive years, plus one additional year if no one else will run."
    )
    body(
        "5. An officer may be removed for good cause by a majority of members present at a general meeting called for "
        "that purpose, with at least 14 days’ notice of time, place, and purpose."
    )
    body(
        "6. If the presidency is vacant, the Vice President becomes President. Other vacancies are filled by majority "
        "vote of the Executive Board for the unexpired term."
    )
    hdg("Article VII — Duties of Officers")
    body(
        "President. Preside at Executive Board and general meetings. Appoint committee chairs when seats are vacant. "
        "Create committees with Executive Board consent. Serve as liaison to the principal and District. Submit the annual "
        "PXU Booster Club Application. Obtain prior written principal approval for fundraising and facility use. Maintain "
        "the spend-approval channel so at least three board members can approve disbursements. Sign contracts, checks, "
        "and disbursements subject to Executive Board approval."
    )
    body(
        "Vice President. Act as President in the President’s absence. Serve as a regular digital approver of expenditures. "
        "Coordinate volunteers. Assist with large events."
    )
    body(
        "Secretary. Record minutes of general and Executive Board meetings and make them available to members. Keep the "
        "membership roll, bylaws, articles, EIN letter, IRS determination, and District approvals. Cross-check spend-approval "
        "logs with minutes. Transfer records to incoming officers within 30 days of election."
    )
    body(
        "Treasurer. Receive, record, and deposit all funds in an account held solely in the Organization’s name. Require "
        "two signatures on every check. Pay only against invoices or original receipts. Reconcile each bank statement within "
        "30 days; a second officer who does not handle money shall review and sign the reconciliation. Keep a digital archive "
        "of approvals paired with receipts. Report at each general meeting. Prepare the year-end financial statement for the "
        "principal and, after tax-exempt recognition, the required IRS Form 990 series. Coordinate receipt backup with the "
        "program sponsor; the sponsor is not a signer."
    )
    hdg("Article VIII — Executive Board duties")
    body(
        "Transact business between general meetings; create committees; approve committee plans; report to the membership; "
        "prepare and approve an annual budget; review financial records. Program sponsors counsel the board on school and "
        "District policy, student needs, and anticipated program costs. They do not vote on expenditures and do not sign checks."
    )
    hdg("Article IX — Meetings")
    body("1. At least two general membership meetings each business year: one in August or September and one in May.")
    body("2. Special general meetings may be called with at least two weeks’ notice of date, time, place, and purpose.")
    body("3. A quorum to transact business is a majority of the Executive Board (not counting ex officio advisors).")
    body(
        "4. Robert’s Rules of Order do not apply. Action is by a seconded motion and majority vote of those entitled to "
        "vote who are present."
    )
    body("5. Disruptive conduct may result in removal from the meeting or from future Organization activities.")
    hdg("Article X — Committees")
    body(
        "The President, with Executive Board consent, may create committees. Committee terms follow the business year. "
        "Each chair presents an annual report and inventory of Organization property at year end."
    )
    hdg("Article XI — Financial controls")
    body("1. Two signatures are required on all checks. Blank checks shall not be issued.")
    body(
        "2. Disbursements require prior approval of at least three Executive Board members (or a minuted board vote), "
        "except that the board may set a written petty-cash limit."
    )
    body("3. The Organization shall not maintain individual student accounts or credit fundraising to a particular student.")
    body("4. The Organization shall not use the District’s taxpayer identification number or school letterhead.")
    body("5. Cash from events shall be counted by two members, documented, and deposited promptly.")
    body(
        "6. Financial records shall be kept five years. An independent third party (not required to be a CPA) should review the books annually."
    )
    hdg("Article XII — Amendments")
    body(
        "These bylaws may be amended by majority vote of members present and voting at a meeting called for that purpose, "
        "provided the amendment is consistent with Arizona law, PXU policy, and section 501(c)(3)."
    )
    body("Adopted this ____ day of ______________, 20____.")
    body("President: ______________________________    Signature: ______________________________    Date: __________")
    body("Vice President: __________________________    Signature: ______________________________    Date: __________")
    body("Secretary: ______________________________    Signature: ______________________________    Date: __________")
    body("Treasurer: ______________________________    Signature: ______________________________    Date: __________")
    doc.save(BYLAWS)


def write_roster_docx() -> None:
    doc = Document()
    t = doc.add_paragraph("Camelback High School Booster Club — Officer & Filing Roster")
    t.runs[0].bold = True
    t.runs[0].font.size = Pt(14)
    doc.add_paragraph(
        "Fill this, then copy onto the PXU Booster Club Application, Arizona articles, EIN application, "
        "and Form 1023-EZ. Do not write Social Security numbers on this sheet. The EIN responsible party "
        "enters an SSN only on IRS.gov."
    )
    meta = doc.add_table(rows=6, cols=2)
    meta.style = "Table Grid"
    labels = [
        ("Working / legal name (principal-approved)", ""),
        ("Program supported (JROTC, band, sport, etc.)", ""),
        ("Mailing address (not the school mailroom)", ""),
        ("Arizona street address (known place of business)", ""),
        ("Fiscal year end", "July 31 (default)"),
        ("Mission / purpose (1–2 sentences)", ""),
    ]
    for i, (lab, val) in enumerate(labels):
        meta.rows[i].cells[0].text = lab
        meta.rows[i].cells[1].text = val

    doc.add_paragraph("")
    doc.add_paragraph("People (officers, agent, sponsor, extra directors/members)").runs[0].bold = True
    people = doc.add_table(rows=1, cols=6)
    people.style = "Table Grid"
    headers = ["Role", "Full legal name", "Mailing address", "Phone", "Email", "Notes (director? incorporator? signer?)"]
    for i, h in enumerate(headers):
        people.rows[0].cells[i].text = h
    roles = [
        "President",
        "Vice President",
        "Secretary",
        "Treasurer",
        "Statutory agent",
        "School sponsor (advisor only — not an officer)",
        "Director / member",
        "Director / member",
        "Director / member",
        "EIN responsible party (name/title only — no SSN here)",
        "Check signer 1",
        "Check signer 2",
    ]
    for role in roles:
        row = people.add_row().cells
        row[0].text = role

    doc.add_paragraph("")
    doc.add_paragraph(
        "PXU reminder: district employees cannot be officers or check signers. Two signatures on every check. "
        "Do not use the school name or PUHSD in the club title or bank account without principal permission."
    )
    doc.save(ROSTER)


def main() -> int:
    KIT.mkdir(parents=True, exist_ok=True)
    spec = build_spec()
    SPEC.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    cmd = [sys.executable, str(PDF_CREATE), str(SPEC), "-o", str(PDF)]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    sys.stdout.write(proc.stdout)
    sys.stderr.write(proc.stderr)
    if proc.returncode != 0:
        return proc.returncode
    write_bylaws_docx()
    write_roster_docx()
    print("WROTE", PDF)
    print("WROTE", BYLAWS)
    print("WROTE", ROSTER)
    print("PDF_BYTES", PDF.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
